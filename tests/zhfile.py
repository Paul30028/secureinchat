# W9: Chinese file metadata remover page (/zh/tools/file-metadata-remover/), served with the CSP from
# _headers for that path. Checks the Chinese interface, a real clean of a JPEG and a DOCX, the network
# panel, and that the English page still reads exactly as before.
import http.server, threading, socketserver, os, re, tempfile, zipfile, sys
from playwright.sync_api import sync_playwright
from PIL import Image
import piexif
from docx import Document

ROOT = os.path.join(os.path.dirname(__file__), '..', 'public')
hdr = open(os.path.join(ROOT, '_headers')).read()
m = re.search(r'/zh/tools/file-metadata-remover/\n\s*Content-Security-Policy: (.*)', hdr)
assert m, 'zh file page has no CSP in _headers'
csp_tools = re.search(r'/tools/\*\n\s*Content-Security-Policy: (.*)', hdr).group(1)
assert m.group(1) == csp_tools, 'zh tool page must use the tools CSP'
assert "'wasm-unsafe-eval'" in csp_tools, 'tools CSP must allow WebAssembly compilation'

class H(http.server.SimpleHTTPRequestHandler):
    def __init__(s, *a, **k): super().__init__(*a, directory=ROOT, **k)
    def end_headers(s): s.send_header('Content-Security-Policy', csp_tools); super().end_headers()
    def log_message(s, *a): pass

srv = socketserver.TCPServer(('127.0.0.1', 0), H); port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
base = f'http://127.0.0.1:{port}'
fails = []
def ok(c, msg):
    print(('ok   ' if c else 'FAIL ') + msg)
    if not c: fails.append(msg)

def wait_ready(pg):
    pg.wait_for_function("() => { var b = document.getElementById('clean'); return b && !b.hidden && !b.disabled; }")

def wait_summary(pg, text):
    pg.wait_for_function("(t) => document.getElementById('summary').textContent.indexOf(t) !== -1", arg=text)

shots = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix='zhfile-shots-')
os.makedirs(shots, exist_ok=True)
tmp = tempfile.mkdtemp(prefix='zhfile-')
jpg = os.path.join(tmp, 'gps.jpg')
Image.new('RGB', (64, 48), (120, 30, 200)).save(jpg, 'JPEG', quality=90)
exif = {'0th': {piexif.ImageIFD.Make: b'TestCam', piexif.ImageIFD.Orientation: 1},
        'Exif': {piexif.ExifIFD.DateTimeOriginal: b'2026:10:20 10:00:00'},
        'GPS': {piexif.GPSIFD.GPSLatitudeRef: b'N', piexif.GPSIFD.GPSLatitude: ((37, 1), (46, 1), (3000, 100)),
                piexif.GPSIFD.GPSLongitudeRef: b'W', piexif.GPSIFD.GPSLongitude: ((122, 1), (25, 1), (1000, 100))},
        '1st': {}, 'thumbnail': None}
piexif.insert(piexif.dump(exif), jpg)
docx_src = os.path.join(tmp, '作者.docx')
marker = 'W9MARKER-AUTHOR-41c2'
d = Document(); d.add_paragraph('元数据测试'); d.core_properties.author = marker; d.core_properties.last_modified_by = marker; d.save(docx_src)
txt = os.path.join(tmp, 'note.txt'); open(txt, 'w', encoding='utf-8').write('plain text\n')
doc_old = os.path.join(tmp, 'old.doc')
# OLE2 compound file signature: the engine reports it as an old Office file.
open(doc_old, 'wb').write(bytes.fromhex('D0CF11E0A1B11AE1') + b'\0' * 1024)

with sync_playwright() as p:
    b = p.chromium.launch()
    for width in (1280, 390):
        pg = b.new_page(viewport={'width': width, 'height': 800}, accept_downloads=True)
        errs, reqs = [], []
        pg.on('console', lambda msg: msg.type == 'error' and errs.append(msg.text))
        pg.on('pageerror', lambda e: errs.append(str(e)))
        pg.on('request', lambda r: reqs.append(r.url))
        pg.goto(base + '/zh/tools/file-metadata-remover/'); pg.wait_for_load_state('networkidle')
        tag = f'[{width}] '
        if width == 1280:
            ok(pg.evaluate('document.documentElement.lang') == 'zh-CN', 'lang is zh-CN')
            ok(pg.locator('h1').inner_text().strip() == '文件元数据清理', 'h1')
            ok(pg.title().startswith('文件元数据清理'), 'title in Chinese')
            alts = pg.eval_on_selector_all('link[rel=alternate][hreflang]', 'els => els.map(e => e.hreflang + " " + e.getAttribute("href"))')
            ok(any(a.startswith('en ') and a.endswith('/tools/file-metadata-remover/') for a in alts) and any(a.startswith('zh-CN ') and a.endswith('/zh/tools/file-metadata-remover/') for a in alts), f'hreflang both ways {alts}')
            ok(pg.locator('link[rel=canonical]').get_attribute('href').endswith('/zh/tools/file-metadata-remover/'), 'canonical')
            crumbs = pg.locator('nav.breadcrumbs a')
            ok(crumbs.nth(0).get_attribute('href') == '/zh/' and crumbs.nth(0).inner_text() == '首页', 'breadcrumb home -> /zh/')
            ok(any(h == '/zh/tools/' for h in crumbs.evaluate_all('els => els.map(e => e.getAttribute("href"))')), 'breadcrumb tools -> /zh/tools/')
            side = pg.locator('.tool-sidebar').inner_text()
            ok(all(k in side for k in ['你的数据如何处理', '留在本机：你的文件、读出的元数据和清理后的副本', '本地处理']), 'sidebar data flow in Chinese')
            ok('本页发出的网络请求' in pg.locator('#netlog summary').inner_text(), 'network panel title in Chinese')
            ok(pg.locator('#netlog thead th').all_text_contents() == ['时间', '类型', '地址', '大小'], 'network table headers')
            ok(pg.locator('a[href="/zh/tools/photo-metadata-remover/"]').count() >= 1, 'links the Chinese photo page')
            ok(pg.locator('a[href="/tools/ip-leak-test/"]').count() == 0, 'no link to the IP leak test (T8 rule)')
            ld = pg.eval_on_selector_all('script[type="application/ld+json"]', 'els => els.map(e => e.textContent).join("\\n")')
            ok('"inLanguage": "zh-CN"' in ld and '文件元数据清理' in ld, 'JSON-LD in Chinese')
        ok(pg.evaluate("() => document.getElementById('clean').hidden"), tag + 'clean button hidden before a file')
        ok(pg.locator('#summary').inner_text().strip() == '选择一个文件开始。', tag + 'idle summary in Chinese')
        ok(not any(u.endswith('metadata_remover_bg.wasm') for u in reqs), tag + 'wasm not requested before a file is chosen')

        pg.set_input_files('#file', jpg)
        wait_ready(pg)
        s = pg.locator('#summary').inner_text().strip()
        ok(re.fullmatch(r'在这个 JPEG 文件里找到 \d+ 项元数据。分享前先清理。', s) is not None, tag + f'found summary in Chinese: {s}')
        rows = pg.locator('#rows').inner_text()
        ok('文件' in rows and '信息' in rows and '已找到' in rows, tag + 'row labels in Chinese')
        ok('gps' in rows.lower(), tag + 'GPS group shown')
        ok(pg.locator('#rows th').first.get_attribute('data-label') == '检查', tag + 'mobile cell labels in Chinese')
        ok(pg.locator('#clean').inner_text().strip() == '删除元数据并下载', tag + 'clean button label')
        ok(pg.locator('#netlog .netlog-marker').text_content().strip() == '你选择了一个文件', tag + 'marker row in Chinese')
        ok(any(u.endswith('metadata_remover_bg.wasm') for u in reqs), tag + 'wasm loaded from this site on first file')
        with pg.expect_download() as info:
            pg.click('#clean')
        out = os.path.join(tmp, f'gps-clean-{width}.jpg'); info.value.save_as(out)
        ok(info.value.suggested_filename == 'gps-clean.jpg', tag + 'download name ' + info.value.suggested_filename)
        wait_summary(pg, '已清理')
        ok(pg.locator('#summary').inner_text().strip() == '已清理。再次扫描没有发现剩余的元数据。清理后的副本已保存到下载文件夹，原文件没有改动。', tag + 'clean message')
        try: gps_after = piexif.load(out).get('GPS') or {}
        except Exception: gps_after = {}
        ok(not gps_after, tag + 'downloaded JPEG has no GPS')

        if width == 1280:
            pg.set_input_files('#file', docx_src)
            wait_ready(pg)
            with pg.expect_download() as info2:
                pg.click('#clean')
            out2 = os.path.join(tmp, 'author-clean.docx'); info2.value.save_as(out2)
            ok(info2.value.suggested_filename == '作者-clean.docx', 'DOCX download name ' + info2.value.suggested_filename)
            with zipfile.ZipFile(out2) as z:
                leaks = [n for n in z.namelist() if marker.encode() in z.read(n)]
            ok(not leaks, f'cleaned DOCX has no author marker {leaks}')
            Document(out2)
            ok('已保存' in pg.locator('#summary').inner_text(), 'DOCX summary says saved')

            pg.set_input_files('#file', txt)
            wait_summary(pg, '不支持')
            ok(pg.locator('#summary').inner_text().strip().startswith('不支持这种文件类型。支持的格式：JPEG、PNG'), 'unsupported message in Chinese')
            ok(pg.evaluate("() => document.getElementById('clean').hidden"), 'clean button hidden for an unsupported file')
            pg.set_input_files('#file', doc_old)
            wait_summary(pg, '旧版 Office')
            ok('另存为 .docx' in pg.locator('#summary').inner_text(), 'old Office message in Chinese')

        ok(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), tag + 'no horizontal scroll')
        ext = [u for u in reqs if not (u.startswith(base) or u.startswith('blob:') or u.startswith('data:'))]
        ok(not ext, tag + f'no outside requests {ext}')
        ok(not errs, tag + f'no console or CSP errors {errs}')
        pg.screenshot(path=os.path.join(shots, f'zhfile-{width}.png'), full_page=True)
        pg.close()

    # English page: same script, old strings.
    pg = b.new_page()
    pg.goto(base + '/tools/file-metadata-remover/'); pg.wait_for_load_state('networkidle')
    ok(pg.evaluate('document.documentElement.lang') == 'en', 'English page lang en')
    alts = pg.eval_on_selector_all('link[rel=alternate][hreflang]', 'els => els.map(e => e.hreflang)')
    ok('zh-CN' in alts, 'English page links its Chinese version')
    pg.set_input_files('#file', jpg)
    wait_ready(pg)
    ok(re.fullmatch(r'\d+ metadata items? found in this JPEG file\. Clean it before you share it\.', pg.locator('#summary').inner_text().strip()) is not None, 'English summary unchanged')
    ok(pg.locator('#netlog .netlog-marker').text_content().strip() == 'You picked a file', 'English marker unchanged')
    ok(pg.locator('#rows th').first.get_attribute('data-label') == 'Check', 'English cell labels unchanged')
    b.close()

idx = open(os.path.join(ROOT, 'zh/tools/index.html'), encoding='utf-8').read()
ok('href="/zh/tools/file-metadata-remover/"' in idx and 'href="/tools/file-metadata-remover/"' not in idx, '/zh/tools/ links the Chinese page')
ok('/zh/tools/file-metadata-remover/' in open(os.path.join(ROOT, 'sitemap.xml')).read(), 'in sitemap')
srv.shutdown()
print('screenshots:', shots)
print('FAILED: %d' % len(fails) if fails else 'all passed')
raise SystemExit(1 if fails else 0)
