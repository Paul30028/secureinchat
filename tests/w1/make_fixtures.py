# Builds files that carry known metadata, for the W1 evaluation of iAnonymous3000/metadata-remover.
# Each fixture has a marker string (or GPS) we can look for after cleaning.
# Usage: python3 tests/w1/make_fixtures.py OUTDIR
import io, os, sys, zipfile
from PIL import Image, PngImagePlugin
import piexif
from pypdf import PdfWriter
import docx

out = sys.argv[1]; os.makedirs(out, exist_ok=True)
MARK = 'W1MARKER-Paul-Test'

def gps():
    def r(x): return (int(x * 1000000), 1000000)
    return {piexif.GPSIFD.GPSLatitudeRef: b'N', piexif.GPSIFD.GPSLatitude: ((31, 1), (13, 1), r(48.5)),
            piexif.GPSIFD.GPSLongitudeRef: b'E', piexif.GPSIFD.GPSLongitude: ((121, 1), (28, 1), r(12.25))}

def exif_bytes():
    return piexif.dump({'0th': {piexif.ImageIFD.Make: b'TestMake', piexif.ImageIFD.Model: b'TestModel',
                                piexif.ImageIFD.Artist: MARK.encode(), piexif.ImageIFD.Software: b'W1Soft'},
                        'Exif': {piexif.ExifIFD.DateTimeOriginal: b'2026:10:28 09:00:00',
                                 piexif.ExifIFD.UserComment: b'ASCII\0\0\0' + MARK.encode()},
                        'GPS': gps()})

img = Image.new('RGB', (640, 480), (200, 120, 40))
for y in range(0, 480, 8):
    for x in range(0, 640, 8): img.putpixel((x, y), (x % 255, y % 255, 90))
img.save(f'{out}/photo.jpg', 'JPEG', quality=90, exif=exif_bytes())
img.rotate(90, expand=True).save(f'{out}/photo-rotated.jpg', 'JPEG', quality=90,
    exif=piexif.dump({'0th': {piexif.ImageIFD.Orientation: 6, piexif.ImageIFD.Artist: MARK.encode()}, 'GPS': gps()}))
# A large photo for timing (about 12 MP).
big = Image.effect_noise((4000, 3000), 64).convert('RGB')
big.save(f'{out}/large.jpg', 'JPEG', quality=92, exif=exif_bytes())

meta = PngImagePlugin.PngInfo(); meta.add_text('Author', MARK); meta.add_text('Comment', MARK)
img.save(f'{out}/image.png', 'PNG', pnginfo=meta, exif=exif_bytes())
img.save(f'{out}/image.webp', 'WEBP', exif=exif_bytes())

w = PdfWriter(); w.add_blank_page(width=200, height=200)
w.add_metadata({'/Author': MARK, '/Title': MARK, '/Creator': 'W1Creator', '/Producer': 'W1Producer'})
w.write(f'{out}/doc.pdf')

d = docx.Document(); d.add_paragraph('Body text stays.')
cp = d.core_properties; cp.author = MARK; cp.last_modified_by = MARK; cp.title = MARK; cp.comments = MARK
d.save(f'{out}/doc.docx')

# A DOCX with an embedded JPEG that itself has GPS (the README says embedded images are cleaned too).
d = docx.Document(); d.add_paragraph('Has a photo.'); d.add_picture(f'{out}/photo.jpg')
d.core_properties.author = MARK; d.save(f'{out}/doc-with-photo.docx')
# HEIC (what iPhones save) with EXIF and GPS, if pillow-heif is installed.
try:
    import pillow_heif
    pillow_heif.register_heif_opener()
    img.save(f'{out}/photo.heic', 'HEIF', exif=exif_bytes())
except ImportError:
    print('pillow-heif not installed: no HEIC fixture')
# MP4 with a title, artist and an ISO 6709 location tag (as phones write), if ffmpeg is installed.
import shutil, subprocess
if shutil.which('ffmpeg'):
    subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=10', '-t', '2',
                    '-pix_fmt', 'yuv420p', '-metadata', f'title={MARK}', '-metadata', f'artist={MARK}',
                    '-metadata', 'location=+31.2304+121.4737/', '-movflags', '+use_metadata_tags', f'{out}/video.mp4'], check=True)
else:
    print('ffmpeg not installed: no MP4 fixture')
print('fixtures in', out, sorted(os.listdir(out)))
