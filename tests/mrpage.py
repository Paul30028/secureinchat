import http.server, threading, socketserver, os, re, tempfile, zipfile, shutil
from playwright.sync_api import sync_playwright
from PIL import Image
import piexif
from docx import Document

ROOT = os.path.join(os.path.dirname(__file__), '..', 'public')
csp = re.search(r'/tools/\*\n\s*Content-Security-Policy: (.*)', open(os.path.join(ROOT, '_headers')).read()).group(1)

class H(http.server.SimpleHTTPRequestHandler):
    def __init__(s, *a, **k): super().__init__(*a, directory=ROOT, **k)
    def end_headers(s): s.send_header('Content-Security-Policy', csp); super().end_headers()
    def log_message(s, *a): pass

srv = socketserver.TCPServer(('127.0.0.1', 0), H); port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
base = f'http://127.0.0.1:{port}'

fails = []
def ok(c, m):
    print(('ok   ' if c else 'FAIL ') + m)
    if not c: fails.append(m)

def wait_ready(pg):
    pg.wait_for_function("() => { var b = document.getElementById('clean'); return b && !b.hidden && !b.disabled; }")

def wait_summary(pg, text):
    pg.wait_for_function("(t) => document.getElementById('summary').textContent.indexOf(t) !== -1", arg=text)

tmp = tempfile.mkdtemp(prefix='mrpage-')
try:
    # --- sample files ---
    jpg = os.path.join(tmp, 'gps.jpg')
    Image.new('RGB', (64, 48), (120, 30, 200)).save(jpg, 'JPEG', quality=90)
    exif = {
        '0th': {piexif.ImageIFD.Make: b'TestCam', piexif.ImageIFD.Orientation: 1},
        'Exif': {piexif.ExifIFD.DateTimeOriginal: b'2026:10:20 10:00:00'},
        'GPS': {
            piexif.GPSIFD.GPSLatitudeRef: b'N',
            piexif.GPSIFD.GPSLatitude: ((37, 1), (46, 1), (3000, 100)),
            piexif.GPSIFD.GPSLongitudeRef: b'W',
            piexif.GPSIFD.GPSLongitude: ((122, 1), (25, 1), (1000, 100)),
        },
        '1st': {}, 'thumbnail': None,
    }
    piexif.insert(piexif.dump(exif), jpg)
    ok(bool(piexif.load(jpg).get('GPS')), 'sample JPEG carries GPS before the test')

    docx_src = os.path.join(tmp, 'author.docx')
    marker = 'W3MARKER-AUTHOR-7f3a'
    d = Document()
    d.add_paragraph('Hello from the metadata test.')
    d.core_properties.author = marker
    d.core_properties.last_modified_by = marker
    d.save(docx_src)
    with zipfile.ZipFile(docx_src) as z:
        ok(any(marker.encode() in z.read(n) for n in z.namelist()), 'sample DOCX carries the marker before the test')

    txt = os.path.join(tmp, 'note.txt')
    with open(txt, 'w', encoding='utf-8') as f:
        f.write('plain text\n')

    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page(viewport={'width': 1280, 'height': 800})
        errs = []; all_reqs = []
        pg.on('pageerror', lambda e: errs.append(str(e)))
        pg.on('console', lambda m: errs.append(m.text) if 'Content Security' in m.text else None)
        pg.on('request', lambda r: all_reqs.append(r.url))
        pg.goto(f'{base}/tools/file-metadata-remover/')

        ok(pg.evaluate("() => document.getElementById('clean').hidden"), 'clean button hidden before a file is chosen')
        ok('Choose a file' in pg.inner_text('#summary'), 'idle summary shown')
        ok(not any(u.endswith('metadata_remover_bg.wasm') for u in all_reqs), 'wasm not requested before a file is chosen')

        # --- JPEG with GPS ---
        pg.set_input_files('#file', jpg)
        wait_ready(pg)
        rows_text = pg.inner_text('#rows').lower()
        ok('gps' in rows_text or 'location' in rows_text, 'GPS row shown for the JPEG')
        ok(any(u.endswith('metadata_remover_bg.wasm') for u in all_reqs), 'wasm requested lazily on first file')

        with pg.expect_download() as info:
            pg.click('#clean')
        dl = info.value
        clean_jpg = os.path.join(tmp, 'gps-clean.jpg')
        dl.save_as(clean_jpg)
        ok(dl.suggested_filename == 'gps-clean.jpg', 'download named <name>-clean.<ext>: ' + dl.suggested_filename)
        wait_summary(pg, 'second scan found no metadata')
        ok('second scan found no metadata' in pg.inner_text('#summary'), 'summary confirms the second scan is clean')
        try:
            gps_after = piexif.load(clean_jpg).get('GPS') or {}
        except Exception:
            gps_after = {}
        ok(not gps_after, 'downloaded JPEG has no GPS left')
        im = Image.open(clean_jpg); im.load()
        ok(im.size == (64, 48), 'downloaded JPEG still opens with the same size')

        # --- DOCX with an author marker ---
        pg.set_input_files('#file', docx_src)
        wait_ready(pg)
        with pg.expect_download() as info2:
            pg.click('#clean')
        dl2 = info2.value
        clean_docx = os.path.join(tmp, 'author-clean.docx')
        dl2.save_as(clean_docx)
        ok(dl2.suggested_filename == 'author-clean.docx', 'DOCX download named <name>-clean.docx')
        with zipfile.ZipFile(clean_docx) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        leaks = [n for n, data in parts.items() if marker.encode() in data]
        ok(not leaks, 'no part of the cleaned DOCX contains the marker ' + str(leaks))
        Document(clean_docx)
        ok(True, 'cleaned DOCX still opens with python-docx')
        ok('saved' in pg.inner_text('#summary'), 'DOCX summary says the copy was saved')

        # --- unsupported .txt ---
        pg.set_input_files('#file', txt)
        wait_summary(pg, 'not supported')
        ok('not supported' in pg.inner_text('#summary'), 'txt reports an unsupported type')
        ok(pg.evaluate("() => document.getElementById('clean').hidden"), 'clean button hidden for unsupported file')

        # --- global checks ---
        external = [u for u in all_reqs if not u.startswith(base) and not u.startswith('blob:') and not u.startswith('data:')]
        ok(not external, 'no external requests ' + str(external))
        ok(not errs, 'no page or CSP errors ' + str(errs))
        pg.set_viewport_size({'width': 390, 'height': 800})
        ok(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'no overflow at 390px')
        b.close()
finally:
    shutil.rmtree(tmp, ignore_errors=True)

print('MR PAGE ' + ('FAILED' if fails else 'PASSED'))
raise SystemExit(1 if fails else 0)
